import os
import ssl
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.template.loader import render_to_string
from django.db.models import Sum, Q
from django.db import transaction
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa

# Bypass SSL certificate verification for Cloudinary image fetching in xhtml2pdf on Windows
try:
    ssl._create_default_https_context = ssl._create_unverified_context
except AttributeError:
    pass

import json
from datetime import date, timedelta, datetime
from django.utils import timezone
from django.db.models import Sum, Count, Q
from django.db.models.functions import Coalesce

from SIA.models import SIA_producto, EntradaInventario, DetalleEntradaInventario
from entregas.models import BeneficioEntregado
from core.models import RegistroAuditoria
from core.utils import registrar_auditoria, obtener_url_inicio_usuario, parse_date_safe

def obtener_datos_estadisticas_entregas(filtro='7d', fecha_inicio_str=None, fecha_fin_str=None, status_filtro='Entregado'):
    """
    Calcula estadísticas agregadas y series temporales diarias de entregas
    para el Dashboard según el filtro seleccionado (hoy, 7d, 30d, mes, personalizado, todos).
    Filtra y contabiliza de manera exclusiva los registros con estatus 'Entregado'.
    """
    hoy = timezone.localdate() if hasattr(timezone, 'localdate') else date.today()
    
    # Filtro estricto: solo beneficios con estatus 'Entregado'
    # Usamos Coalesce para tomar fecha_entrega si existe, de lo contrario fecha de registro
    qs = BeneficioEntregado.objects.filter(status__iexact='Entregado').annotate(
        fecha_efectiva=Coalesce('fecha_entrega', 'fecha')
    )

    fecha_desde = None
    fecha_hasta = hoy
    chart_tipo = 'line'
    chart_subtitulo = ''

    if filtro == 'hoy':
        fecha_desde = hoy
        fecha_hasta = hoy
        qs = qs.filter(Q(fecha_efectiva=hoy) | Q(fecha=hoy) | Q(fecha_entrega=hoy))
    elif filtro == '7d':
        fecha_desde = hoy - timedelta(days=6)
        fecha_hasta = hoy
        qs = qs.filter(
            Q(fecha_efectiva__gte=fecha_desde, fecha_efectiva__lte=fecha_hasta) |
            Q(fecha__gte=fecha_desde, fecha__lte=fecha_hasta)
        )
    elif filtro == '30d':
        fecha_desde = hoy - timedelta(days=29)
        fecha_hasta = hoy
        qs = qs.filter(
            Q(fecha_efectiva__gte=fecha_desde, fecha_efectiva__lte=fecha_hasta) |
            Q(fecha__gte=fecha_desde, fecha__lte=fecha_hasta)
        )
    elif filtro == 'mes':
        fecha_desde = date(hoy.year, hoy.month, 1)
        fecha_hasta = hoy
        qs = qs.filter(
            Q(fecha_efectiva__year=hoy.year, fecha_efectiva__month=hoy.month) |
            Q(fecha__year=hoy.year, fecha__month=hoy.month)
        )
    elif filtro == 'personalizado':
        q_filt = Q()
        if fecha_inicio_str:
            d_ini = parse_date_safe(fecha_inicio_str)
            if d_ini:
                fecha_desde = d_ini
                q_filt &= (Q(fecha_efectiva__gte=d_ini) | Q(fecha__gte=d_ini))
        if fecha_fin_str:
            d_fin = parse_date_safe(fecha_fin_str)
            if d_fin:
                fecha_hasta = d_fin
                q_filt &= (Q(fecha_efectiva__lte=d_fin) | Q(fecha__lte=d_fin))
        if q_filt:
            qs = qs.filter(q_filt)
    elif filtro == 'todos':
        fecha_desde = None
        fecha_hasta = None

    # Métricas KPIs del período seleccionado
    total_entregas = qs.count()
    total_articulos = qs.aggregate(total=Sum('cantidad_dada'))['total'] or 0
    total_beneficiarios = qs.exclude(cedula__isnull=True).exclude(cedula__exact='').values('cedula').distinct().count()

    # Cálculo de días para promedio de entregas
    if filtro == 'hoy':
        dias_conteo = 1
    elif fecha_desde and fecha_hasta:
        dias_conteo = max(1, (fecha_hasta - fecha_desde).days + 1)
    else:
        dias_activos = qs.filter(fecha_efectiva__isnull=False).values('fecha_efectiva').distinct().count()
        dias_conteo = max(1, dias_activos)

    promedio_diario = round(total_entregas / dias_conteo, 1) if dias_conteo > 0 else 0

    # Top productos entregados en el período
    top_prods_qs = qs.values('descripcion_prod').annotate(
        total_cant=Sum('cantidad_dada'),
        total_entregas=Count('id')
    ).order_by('-total_cant')[:6]

    top_productos = [{
        'descripcion': p['descripcion_prod'] or 'Sin descripción',
        'cantidad': int(p['total_cant'] or 0),
        'entregas': p['total_entregas']
    } for p in top_prods_qs]

    chart_labels = []
    chart_entregas = []
    chart_unidades = []

    # Configuración específica para el filtro HOY
    if filtro == 'hoy':
        chart_tipo = 'bar'
        if total_entregas > 0 and top_productos:
            # Si hoy hubo entregas, graficamos los insumos entregados hoy
            chart_subtitulo = f"Entregas de hoy ({hoy.strftime('%d/%m')}) por insumo"
            chart_labels = [p['descripcion'] for p in top_productos]
            chart_entregas = [p['entregas'] for p in top_productos]
            chart_unidades = [p['cantidad'] for p in top_productos]
        else:
            # Si hoy aún no hay entregas, mostramos la actividad de los últimos 7 días terminando en Hoy
            # para que el gráfico NUNCA quede vacío ni desaparezca
            chart_subtitulo = f"Actividad reciente (Hoy: 0 entregas)"
            curr = hoy - timedelta(days=6)
            entregas_recientes = BeneficioEntregado.objects.filter(
                status__iexact='Entregado'
            ).annotate(f_op=Coalesce('fecha_entrega', 'fecha')).filter(
                f_op__gte=curr, f_op__lte=hoy
            ).values('f_op').annotate(
                cnt=Count('id'), cant=Sum('cantidad_dada')
            )
            dict_recientes = {item['f_op']: item for item in entregas_recientes}
            while curr <= hoy:
                lbl = curr.strftime('%d/%m') + (' (Hoy)' if curr == hoy else '')
                chart_labels.append(lbl)
                if curr in dict_recientes:
                    chart_entregas.append(dict_recientes[curr]['cnt'])
                    chart_unidades.append(int(dict_recientes[curr]['cant'] or 0))
                else:
                    chart_entregas.append(0)
                    chart_unidades.append(0)
                curr += timedelta(days=1)
    else:
        # Agrupación por fecha para evolución temporal (7d, 30d, mes, personalizado, todos)
        agrupado = qs.filter(fecha_efectiva__isnull=False).values('fecha_efectiva').annotate(
            entregas_cnt=Count('id'),
            unidades_cnt=Sum('cantidad_dada')
        ).order_by('fecha_efectiva')

        dict_fechas = {item['fecha_efectiva']: item for item in agrupado}

        if filtro in ['7d', '30d', 'mes'] and fecha_desde and fecha_hasta:
            curr = fecha_desde
            while curr <= fecha_hasta:
                lbl = curr.strftime('%d/%m')
                chart_labels.append(lbl)
                if curr in dict_fechas:
                    chart_entregas.append(dict_fechas[curr]['entregas_cnt'])
                    chart_unidades.append(int(dict_fechas[curr]['unidades_cnt'] or 0))
                else:
                    chart_entregas.append(0)
                    chart_unidades.append(0)
                curr += timedelta(days=1)
            chart_tipo = 'line'
        else:
            for item in agrupado:
                f = item['fecha_efectiva']
                lbl = f.strftime('%d/%m/%y') if filtro == 'todos' else f.strftime('%d/%m')
                chart_labels.append(lbl)
                chart_entregas.append(item['entregas_cnt'])
                chart_unidades.append(int(item['unidades_cnt'] or 0))
            chart_tipo = 'bar' if len(chart_labels) <= 2 else 'line'

    return {
        'filtro': filtro,
        'status_filtro': status_filtro,
        'fecha_desde': fecha_desde.strftime('%Y-%m-%d') if fecha_desde else '',
        'fecha_hasta': fecha_hasta.strftime('%Y-%m-%d') if fecha_hasta else '',
        'total_entregas': total_entregas,
        'total_articulos': total_articulos,
        'total_beneficiarios': total_beneficiarios,
        'promedio_diario': promedio_diario,
        'chart_tipo': chart_tipo,
        'chart_subtitulo': chart_subtitulo,
        'chart_labels': chart_labels,
        'chart_entregas': chart_entregas,
        'chart_unidades': chart_unidades,
        'top_productos': top_productos,
    }

@login_required
def dashboard(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_dashboard:
            target_url = obtener_url_inicio_usuario(request.user)
            if target_url != 'dashboard':
                return redirect(target_url)

    if not request.session.get('bienvenida_mostrada', False):
        messages.success(request, f'¡Hola, {request.user.username}! Bienvenido de nuevo.', extra_tags='bienvenida')
        request.session['bienvenida_mostrada'] = True
    
    entregas = BeneficioEntregado.objects.all()
    pendientes = entregas.filter(status__iexact='Pendiente').count()
    entregados = entregas.filter(status__iexact='Entregado').count()
    total_articulos = SIA_producto.objects.count()
    
    recientes = list(BeneficioEntregado.objects.all().order_by('-id')[:5].values(
        'nombre_beneficiario', 'descripcion_prod', 'fecha'
    ))
    
    resumen_productos = entregas.values('descripcion_prod').annotate(total_cant=Sum('cantidad_dada')).order_by('-total_cant')
    nombres_productos = [item['descripcion_prod'] or 'Desconocido' for item in resumen_productos]
    cantidades_productos = [int(item['total_cant'] or 0) for item in resumen_productos]

    # Pre-cálculo de estadísticas de entregas (por defecto: últimos 7 días con estatus 'Entregado')
    stats_entregas_init = obtener_datos_estadisticas_entregas(filtro='7d', status_filtro='Entregado')

    context = {
        'pendientes': pendientes,
        'entregados': entregados,
        'total_articulos': total_articulos, 
        'recientes': recientes,
        'nombres_productos': nombres_productos,
        'cantidades_productos': cantidades_productos,
        'stats_entregas_init_json': json.dumps(stats_entregas_init),
    }

    return render(request, 'dashboard.html', context)

@login_required
def api_estadisticas_entregas_dashboard(request):
    """
    Endpoint JSON para filtrar dinámicamente las estadísticas de entregas del Dashboard.
    Filtra y procesa exclusivamente registros con estatus 'Entregado'.
    """
    filtro = request.GET.get('filtro', '7d').strip().lower()
    fecha_inicio = request.GET.get('fecha_inicio', '').strip()
    fecha_fin = request.GET.get('fecha_fin', '').strip()

    datos = obtener_datos_estadisticas_entregas(
        filtro=filtro,
        fecha_inicio_str=fecha_inicio,
        fecha_fin_str=fecha_fin,
        status_filtro='Entregado'
    )
    return JsonResponse(datos)

@login_required
def api_notificaciones_resumen(request):
    alertas_productos = list(SIA_producto.objects.filter(cantidad__lte=5).values('codigo', 'descripcion', 'cantidad'))
    pendientes_count = BeneficioEntregado.objects.filter(status='Pendiente').count()
    
    return JsonResponse({
        'alertas_stock': alertas_productos,
        'alertas_count': len(alertas_productos),
        'pendientes_count': pendientes_count,
        'total_notificaciones': len(alertas_productos) + (1 if pendientes_count > 0 else 0)
    })

@login_required
def generar_reporte_inventario_independiente(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido: No tienes permiso para generar reportes de Inventario.")
            return redirect(obtener_url_inicio_usuario(request.user))
    try:
        productos = SIA_producto.objects.all()
        html = render_to_string('reportes/inventario_pdf.html', {'productos': productos})
        
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'inline; filename="Reporte_Inventario_Actual.pdf"'
        
        pisa_status = pisa.CreatePDF(html, dest=response)
        
        if pisa_status.err:
            return HttpResponse('Error interno al generar el PDF de inventario', status=500)
            
        return response
        
    except Exception as e:
        return HttpResponse(f'Error al generar reporte: {str(e)}', status=500)

@login_required
def exportar_inventario_excel(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido: No tienes permiso para exportar el Inventario.")
            return redirect(obtener_url_inicio_usuario(request.user))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"
    
    ws.merge_cells('A1:E1')
    ws['A1'] = "SIA WEB - REPORTE GENERAL DE INVENTARIO"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = ["Código", "Descripción", "Almacén", "Presentación", "Cantidad Total"]
    ws.append(headers)
    ws.row_dimensions[2].height = 24

    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    productos = SIA_producto.objects.all()
    for row_idx, p in enumerate(productos, start=3):
        presentacion = f"Caja ({p.unidades_por_empaque} u.)" if p.tipo_presentacion == 'caja' else "Unidad"
        ws.append([p.codigo, p.descripcion, p.almacen, presentacion, p.cantidad])
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Reporte_Inventario.xlsx"'
    wb.save(response)
    return response

@login_required
def api_buscar_producto_autocomplete(request):
    q = request.GET.get('q', '').strip()
    if len(q) < 2:
        return JsonResponse({'resultados': []})
    
    productos = SIA_producto.objects.filter(
        Q(codigo__icontains=q) | Q(descripcion__icontains=q)
    )[:10]

    resultados = [{
        'id': p.id,
        'codigo': p.codigo,
        'descripcion': p.descripcion,
        'almacen': p.almacen,
        'cantidad': p.cantidad
    } for p in productos]

    return JsonResponse({'resultados': resultados})

@login_required
def lista_inventario(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido: No tienes permiso para acceder al módulo de Inventario.")
            return redirect(obtener_url_inicio_usuario(request.user))

    productos = SIA_producto.objects.all()
    return render(request, 'inventario.html', {'productos': productos})

def generar_nuevo_codigo_producto():
    ultimo = SIA_producto.objects.order_by('-id').first()
    siguiente_id = (ultimo.id + 1) if ultimo else 1
    codigo = f"PROD-{siguiente_id:04d}"
    while SIA_producto.objects.filter(codigo=codigo).exists():
        siguiente_id += 1
        codigo = f"PROD-{siguiente_id:04d}"
    return codigo

@login_required
@transaction.atomic
def registrar_inventario(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido: No tienes permiso para modificar Inventario.")
            return redirect(obtener_url_inicio_usuario(request.user))
    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip()
        if not codigo:
            codigo = generar_nuevo_codigo_producto()

        descripcion = request.POST.get('descripcion')
        
        try:
            cantidad_ingresada = int(request.POST.get('cantidad', 0))
        except ValueError:
            cantidad_ingresada = 0

        tipo_ingreso = request.POST.get('tipo_ingreso', 'unidad')
        unidades_por_caja = 1
        
        if tipo_ingreso == 'caja':
            try:
                unidades_por_caja = int(request.POST.get('unidades_por_caja', 1))
            except ValueError:
                unidades_por_caja = 1
            
            nueva_cantidad = cantidad_ingresada * unidades_por_caja
        else:
            nueva_cantidad = cantidad_ingresada
        
        producto, creado = SIA_producto.objects.get_or_create(
            codigo=codigo,
            defaults={
                'descripcion': descripcion,
                'almacen': request.POST.get('almacen', 'Principal'),
                'cantidad': nueva_cantidad,
                'tipo_presentacion': tipo_ingreso,
                'unidades_por_empaque': unidades_por_caja
            }
        )
        
        if not creado:
            producto.cantidad += nueva_cantidad
            producto.tipo_presentacion = tipo_ingreso
            producto.unidades_por_empaque = unidades_por_caja
            producto.save()

            registrar_auditoria(
                request,
                f"Actualización de stock para producto '{producto.descripcion}' (Código: {codigo}). Se sumaron {nueva_cantidad} unidades.",
                modulo="Inventario"
            )

            messages.success(request, f"Stock actualizado exitosamente. Total: {producto.cantidad}")
        else:
            registrar_auditoria(
                request,
                f"Registro de nuevo producto '{descripcion}' (Código: {codigo}) con cantidad inicial de {nueva_cantidad}.",
                modulo="Inventario"
            )
            
            messages.success(request, "Producto registrado exitosamente.")
            
        return redirect('inventario')
        
    codigo_sugerido = generar_nuevo_codigo_producto()
    return render(request, 'registrar_inventario.html', {'codigo_sugerido': codigo_sugerido})

@login_required
def verificar_producto(request):
    codigo = request.GET.get('codigo', '')
    try:
        producto = SIA_producto.objects.get(codigo=codigo)
        return JsonResponse({
            'existe': True, 
            'descripcion': producto.descripcion, 
            'almacen': producto.almacen
        })
    except SIA_producto.DoesNotExist:
        return JsonResponse({'existe': False})

@login_required
def obtener_todo_inventario(request):
    productos = list(SIA_producto.objects.values('id', 'codigo', 'descripcion', 'almacen', 'cantidad'))
    return JsonResponse({'productos': productos})

@login_required
@transaction.atomic
def eliminar_producto(request, producto_id):
    producto = get_object_or_404(SIA_producto, codigo=producto_id)
    if request.method == 'POST':
        descripcion_prod = producto.descripcion
        codigo_prod = producto.codigo
        
        producto.delete()

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Eliminación de producto {codigo_prod} ({descripcion_prod}) del almacén.",
            modulo="Inventario"
        )
        messages.success(request, "Producto eliminado exitosamente.")
    return redirect('inventario')

# --- MÓDULO DE RECEPCIÓN Y ENTRADAS DE ALMACÉN ---
@login_required
@transaction.atomic
def registrar_entrada_inventario(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido: No tienes permiso para registrar entradas de almacén.")
            return redirect(obtener_url_inicio_usuario(request.user))

    if request.method == 'POST':
        origen_o_proveedor = request.POST.get('origen_o_proveedor', '').strip()
        nro_guia_o_acta = request.POST.get('nro_guia_o_acta', '').strip()
        observaciones = request.POST.get('observaciones', '').strip()

        productos_ids = request.POST.getlist('producto_id[]')
        cantidades = request.POST.getlist('cantidad_ingresada[]')

        if not origen_o_proveedor or not productos_ids:
            messages.error(request, "Por favor indica el origen/proveedor y añade al menos un producto a la entrada.")
            return redirect('registrar_entrada_inventario')

        entrada = EntradaInventario.objects.create(
            origen_o_proveedor=origen_o_proveedor,
            nro_guia_o_acta=nro_guia_o_acta,
            observaciones=observaciones,
            registrado_por=request.user
        )

        resumen_detalles = []
        for i in range(len(productos_ids)):
            prod_id = productos_ids[i]
            try:
                cant_ingresada = int(cantidades[i])
            except (ValueError, TypeError):
                cant_ingresada = 0

            if cant_ingresada > 0:
                producto = SIA_producto.objects.select_for_update().filter(id=prod_id).first()
                if not producto:
                    continue
                producto.cantidad += cant_ingresada
                producto.save()

                DetalleEntradaInventario.objects.create(
                    entrada=entrada,
                    producto=producto,
                    cantidad_ingresada=cant_ingresada
                )
                resumen_detalles.append(f"{producto.descripcion} (+{cant_ingresada} u.)")

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Recepción de mercancía Entrada #{entrada.id} ({origen_o_proveedor}) - Productos: {', '.join(resumen_detalles)}.",
            modulo="Inventario"
        )

        messages.success(request, f"Entrada de inventario #{entrada.id} registrada correctamente. El stock ha sido actualizado.")
        return redirect('lista_entradas_inventario')

    productos = SIA_producto.objects.all().order_by('descripcion')
    return render(request, 'SIA/registrar_entrada.html', {'productos': productos})

@login_required
def lista_entradas_inventario(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido.")
            return redirect(obtener_url_inicio_usuario(request.user))

    entradas = EntradaInventario.objects.all().prefetch_related('detalles__producto').order_by('-fecha')
    return render(request, 'SIA/lista_entradas.html', {'entradas': entradas})

def ficha_producto_view(request, producto_id):
    producto = get_object_or_404(SIA_producto, id=producto_id)
    
    entregas_recientes = BeneficioEntregado.objects.filter(
        descripcion_prod__icontains=producto.descripcion
    ).order_by('-id')[:5]
    
    entradas_recientes = DetalleEntradaInventario.objects.filter(
        producto=producto
    ).select_related('entrada').order_by('-id')[:5]
    
    return render(request, 'SIA/ficha_producto.html', {
        'producto': producto,
        'codigo_prod': f"PROD-{producto.id:04d}",
        'entregas_recientes': entregas_recientes,
        'entradas_recientes': entradas_recientes
    })

@login_required
def imprimir_etiquetas_producto(request, producto_id):
    producto = get_object_or_404(SIA_producto, id=producto_id)
    
    server_ip = request.META.get('HTTP_HOST', 'localhost:8000')
    scheme = 'https' if request.is_secure() else 'http'
    qr_url = f"{scheme}://{server_ip}/inventario/ficha/{producto.id}/"

    return render(request, 'SIA/etiquetas_producto.html', {
        'producto': producto,
        'qr_url': qr_url,
        'codigo_prod': f"PROD-{producto.id:04d}"
    })

@login_required
def descargar_plantilla_excel_view(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla Inventario"

    headers = ["DESCRIPCION", "CANTIDAD", "ALMACEN", "CATEGORIA", "PRESENTACION"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ejemplos = [
        ["Sillas de Ruedas Ejecutivas", 15, "Almacén Central", "Enseres", "unidad"],
        ["Cajas de Paracetamol 500mg", 50, "Depósito Médico", "Medicamentos", "caja"],
        ["Colchones Ortopédicos Matrimoniales", 10, "Depósito General", "Enseres", "unidad"],
    ]
    for row in ejemplos:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="Plantilla_Carga_Inventario_SIA.xlsx"'
    wb.save(response)
    return response

@login_required
def importar_inventario_excel_view(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_inventario:
            messages.error(request, "Acceso restringido: No tienes permiso para importar inventario.")
            return redirect('inventario')

    if request.method == 'POST' and request.FILES.get('archivo_excel'):
        archivo = request.FILES['archivo_excel']
        if not archivo.name.endswith('.xlsx'):
            messages.error(request, "Formato no válido. Debe seleccionar un archivo de Excel (.xlsx).")
            return redirect('inventario')

        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            ws = wb.active
            
            importados = 0
            actualizados = 0

            with transaction.atomic():
                for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not row or not row[0]:
                        continue

                    descripcion = str(row[0]).strip()
                    try:
                        cantidad = int(row[1]) if len(row) > 1 and row[1] is not None else 1
                    except (ValueError, TypeError):
                        cantidad = 1

                    almacen = str(row[2]).strip() if len(row) > 2 and row[2] else "Almacén Central"
                    categoria = str(row[3]).strip() if len(row) > 3 and row[3] else "General"
                    presentacion = str(row[4]).strip().lower() if len(row) > 4 and row[4] else "unidad"
                    if presentacion not in ['unidad', 'caja']:
                        presentacion = 'unidad'

                    prod, created = SIA_producto.objects.get_or_create(
                        descripcion__iexact=descripcion,
                        defaults={
                            'descripcion': descripcion,
                            'cantidad': cantidad,
                            'almacen': almacen,
                            'categoria': categoria,
                            'tipo_presentacion': presentacion,
                            'codigo': f"PROD-TEMP"
                        }
                    )

                    if created:
                        prod.codigo = f"PROD-{prod.id:04d}"
                        prod.save()
                        importados += 1
                    else:
                        prod.cantidad += cantidad
                        prod.save()
                        actualizados += 1

            RegistroAuditoria.objects.create(
                usuario=request.user,
                accion=f"Carga masiva de inventario realizada ({importados} creados, {actualizados} actualizados).",
                modulo="Inventario"
            )
            messages.success(request, f"¡Importación completada con éxito! Se registraron {importados} nuevos productos y se actualizaron {actualizados} existentes.")
        except Exception as e:
            messages.error(request, f"Error al procesar el archivo Excel: {str(e)}")

    return redirect('inventario')
