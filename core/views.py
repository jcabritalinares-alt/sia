from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.http import JsonResponse, HttpResponse, FileResponse, Http404
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth import login, logout
from django.core.cache import cache
from django.utils import timezone
from django.contrib.auth.models import User
from django.core.management import call_command
from django.conf import settings
import os
import socket
import shutil
import json
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from django.db import transaction
from django.db.models import Q, Count, Sum
from django.core.paginator import Paginator
from django.template.loader import render_to_string
from xhtml2pdf import pisa
from core.models import RegistroAuditoria, SolicitudCiudadano, PerfilUsuario
from beneficiarios.models import Beneficiario
from entregas.models import BeneficioEntregado
from SIA.models import SIA_producto
from core.utils import obtener_url_inicio_usuario, parse_date_safe, enviar_notificacion_solicitud_email, generar_hash_firma_digital, subir_archivo_evidencia_seguro

def obtener_ip_cliente(request):
    """Extrae la dirección IP real del cliente, incluso tras proxies y túneles (Cloudflare, etc.)."""
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        return x_forwarded_for.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR', '127.0.0.1')

def login_view(request):
    ip = obtener_ip_cliente(request)
    lockout_key = f'login_lockout_{ip}'
    attempts_key = f'login_attempts_{ip}'

    # Verificar si la IP se encuentra temporalmente bloqueada por exceso de intentos
    if cache.get(lockout_key):
        messages.error(
            request,
            "⚠️ Demasiados intentos fallidos consecutivos. Por motivos de seguridad, tu acceso ha sido bloqueado temporalmente por 10 minutos."
        )
        return render(request, 'login.html', {'form': AuthenticationForm()})

    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            
            # Limpiar intentos fallidos acumulados
            cache.delete(attempts_key)
            cache.delete(lockout_key)

            cache_key = f'last_seen_{user.id}'
            
            # Iniciar sesión e invalidar cualquier rastro de sesión anterior
            login(request, user)
            cache.set(cache_key, timezone.now(), 900)
            request.session['bienvenida_mostrada'] = False 
            request.session['show_welcome_splash'] = True 
            
            RegistroAuditoria.objects.create(
                usuario=user,
                accion=f"Inicio de sesión exitoso para el usuario '{user.username}' (IP: {ip}).",
                modulo="Usuarios",
                timestamp=timezone.now()
            )
            return redirect(obtener_url_inicio_usuario(user))
        else:
            intentos = cache.get(attempts_key, 0) + 1
            cache.set(attempts_key, intentos, 300) # Ventana de 5 minutos

            username_intento = request.POST.get('username', 'Desconocido').strip()

            if intentos >= 5:
                cache.set(lockout_key, True, 600) # Bloqueo por 10 minutos
                cache.delete(attempts_key)
                RegistroAuditoria.objects.create(
                    usuario=None,
                    accion=f"Alerta de Seguridad: Bloqueo de IP {ip} por 10 minutos tras 5 intentos fallidos (Usuario: '{username_intento}').",
                    modulo="Seguridad",
                    timestamp=timezone.now()
                )
                messages.error(
                    request,
                    "⚠️ Has excedido el límite de 5 intentos fallidos. Tu acceso ha sido bloqueado temporalmente por 10 minutos."
                )
            else:
                restantes = 5 - intentos
                messages.error(
                    request,
                    f"Usuario o contraseña incorrectos. Te quedan {restantes} intento(s) antes del bloqueo de seguridad."
                )
    else:
        form = AuthenticationForm()
    return render(request, 'login.html', {'form': form})

def logout_view(request):
    if request.user.is_authenticated:
        cache_key = f'last_seen_{request.user.id}'
        cache.delete(cache_key)
    logout(request)
    messages.info(request, "Has cerrado sesión correctamente.")
    return redirect('login')

@login_required
def registrar_usuario(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('usuarios'))):
        messages.error(request, "Acceso restringido: Solo los usuarios autorizados pueden gestionar cuentas y permisos.")
        return redirect('dashboard')

    modulos_list = [
        'dashboard', 'estadisticas', 'inventario', 'entradas', 
        'atencion_ciudadano', 'entregas', 'asignaciones_especiales', 
        'usuarios', 'auditoria', 'respaldos'
    ]

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()
        rol = request.POST.get('rol', '').strip() or 'Personalizado'

        if not username or not password:
            messages.error(request, "El nombre de usuario y la contraseña son obligatorios.")
            return redirect('registrar_usuario')

        if User.objects.filter(username=username).exists():
            messages.error(request, f"El nombre de usuario '{username}' ya existe.")
            return redirect('registrar_usuario')

        es_superuser = ('es_superuser' in request.POST) and request.user.is_superuser

        user = User.objects.create_user(username=username, password=password)
        user.is_superuser = es_superuser
        user.is_staff = True
        user.first_name = rol[:30]
        user.save()

        # Construir permisos granulares manuales para los 10 módulos
        hay_modulos_marcados = any(f'permiso_{m}' in request.POST for m in modulos_list)
        detalle = {}
        for m in modulos_list:
            if hay_modulos_marcados:
                activo = (f'permiso_{m}' in request.POST) or es_superuser
            else:
                # Si no se marcaron módulos específicos en creación rápida, asignar básicos
                activo = (m in ['dashboard', 'entregas', 'atencion_ciudadano']) or es_superuser

            editar_default = (activo and m in ['inventario', 'atencion_ciudadano'])
            if m == 'entregas':
                editar_default = False

            detalle[m] = {
                'activo': activo,
                'leer': (f'permiso_{m}_leer' in request.POST) if (f'permiso_{m}_leer' in request.POST) else activo,
                'crear': (f'permiso_{m}_crear' in request.POST) if (f'permiso_{m}_crear' in request.POST) else (activo and m in ['inventario', 'entradas', 'atencion_ciudadano', 'entregas', 'asignaciones_especiales']),
                'editar': (f'permiso_{m}_editar' in request.POST) if (f'permiso_{m}_editar' in request.POST) else (es_superuser or editar_default),
                'cambiar_estatus': (f'permiso_{m}_cambiar_estatus' in request.POST) if (f'permiso_{m}_cambiar_estatus' in request.POST) else (activo and m == 'entregas'),
                'eliminar': (f'permiso_{m}_eliminar' in request.POST) if (f'permiso_{m}_eliminar' in request.POST) else False,
                'imprimir': (f'permiso_{m}_imprimir' in request.POST) if (f'permiso_{m}_imprimir' in request.POST) else activo,
            }

        PerfilUsuario.objects.create(
            user=user,
            rol=rol,
            permiso_dashboard=detalle['dashboard']['activo'],
            permiso_inventario=detalle['inventario']['activo'],
            permiso_entregas=detalle['entregas']['activo'],
            permiso_atencion_ciudadano=detalle['atencion_ciudadano']['activo'],
            permiso_asignaciones_especiales=detalle['asignaciones_especiales']['activo'],
            permiso_usuarios=detalle['usuarios']['activo'],
            permiso_auditoria=detalle['auditoria']['activo'],
            permisos_detalle=detalle
        )

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Creación manual de usuario '{username}' con rol '{rol}' y permisos personalizados asignados.",
            modulo="Usuarios",
            timestamp=timezone.now()
        )

        messages.success(request, f"Usuario '{username}' registrado exitosamente con rol '{rol}'.")
        return redirect('registrar_usuario')
    
    usuarios = User.objects.all().order_by('-date_joined')
    
    lista_usuarios_data = []
    for u in usuarios:
        rol_default = 'Administrador' if u.is_superuser else 'Personalizado'
        perfil, _ = PerfilUsuario.objects.get_or_create(user=u, defaults={
            'rol': rol_default,
            'permiso_dashboard': True,
            'permiso_inventario': u.is_superuser,
            'permiso_entregas': True,
            'permiso_atencion_ciudadano': True,
            'permiso_asignaciones_especiales': u.is_superuser,
            'permiso_usuarios': u.is_superuser,
            'permiso_auditoria': u.is_superuser
        })

        lista_usuarios_data.append({
            'id': u.id,
            'username': u.username,
            'is_superuser': u.is_superuser,
            'is_staff': u.is_staff,
            'is_active': u.is_active,
            'rol': perfil.rol,
            'perfil': perfil,
            'date_joined': u.date_joined,
            'total_modulos': perfil.total_modulos_activos,
        })

    return render(request, 'registrar_usuario.html', {'usuarios': lista_usuarios_data})

@login_required
def cambiar_password_usuario(request, user_id):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('usuarios'))):
        messages.error(request, "Acceso restringido: Solo los Administradores o usuarios autorizados pueden cambiar contraseñas.")
        return redirect('dashboard')

    if request.method == 'POST':
        u = get_object_or_404(User, id=user_id)
        nueva_clave = request.POST.get('nueva_password', '').strip()
        confirmar_clave = request.POST.get('confirmar_password', '').strip()

        if not nueva_clave or len(nueva_clave) < 3:
            messages.error(request, "La nueva contraseña debe tener al menos 3 caracteres.")
            return redirect('registrar_usuario')

        if nueva_clave != confirmar_clave:
            messages.error(request, "Las contraseñas ingresadas no coinciden.")
            return redirect('registrar_usuario')

        u.set_password(nueva_clave)
        u.save()

        # Limpiar caché por si tiene sesión activa
        cache.delete(f'last_seen_{u.id}')

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Cambio de contraseña realizado para el usuario '{u.username}'.",
            modulo="Usuarios",
            timestamp=timezone.now()
        )

        messages.success(request, f"🔑 Contraseña del usuario '{u.username}' actualizada exitosamente.")
    return redirect('registrar_usuario')

@login_required
def actualizar_permisos_usuario(request, user_id):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('usuarios'))):
        messages.error(request, "Acceso restringido: Solo los Administradores o usuarios autorizados pueden modificar permisos de módulos.")
        return redirect('dashboard')

    if request.method == 'POST':
        u = get_object_or_404(User, id=user_id)
        rol_nombre = request.POST.get('rol', '').strip() or 'Personalizado'
        
        # Superusuario de Django: solo modificable por otro superusuario
        if request.user.is_superuser:
            es_superuser_solicitado = ('es_superuser' in request.POST)
            # Evitar auto-bloqueo del administrador actual
            if u.id == request.user.id and not es_superuser_solicitado:
                messages.warning(request, "No puedes quitarte el permiso maestro de Superusuario a ti mismo.")
                es_superuser_solicitado = True
            u.is_superuser = es_superuser_solicitado

        u.is_staff = True
        u.first_name = rol_nombre[:30]
        u.save()

        perfil, _ = PerfilUsuario.objects.get_or_create(user=u)
        perfil.rol = rol_nombre

        # Guardar detalle granular de acciones por cada uno de los 10 módulos sin forzar sobrescritura
        modulos_list = [
            'dashboard', 'estadisticas', 'inventario', 'entradas', 
            'atencion_ciudadano', 'entregas', 'asignaciones_especiales', 
            'usuarios', 'auditoria', 'respaldos'
        ]
        
        detalle = {}
        for m in modulos_list:
            activo = (f'permiso_{m}' in request.POST) or u.is_superuser
            leer_val = (f'permiso_{m}_leer' in request.POST) if activo else False
            crear_val = (f'permiso_{m}_crear' in request.POST) if activo else False
            editar_val = (f'permiso_{m}_editar' in request.POST) if activo else False
            cambiar_estatus_val = (f'permiso_{m}_cambiar_estatus' in request.POST) if activo else False
            eliminar_val = (f'permiso_{m}_eliminar' in request.POST) if activo else False
            imprimir_val = (f'permiso_{m}_imprimir' in request.POST) if activo else False

            # Si el módulo está activo y el admin marcó alguna acción pero no 'leer', asegurar 'leer'
            if activo and not leer_val and (crear_val or editar_val or cambiar_estatus_val or imprimir_val):
                leer_val = True

            detalle[m] = {
                'activo': activo,
                'leer': leer_val or (activo and m in ['dashboard', 'auditoria']),
                'crear': crear_val,
                'editar': editar_val,
                'cambiar_estatus': cambiar_estatus_val,
                'eliminar': eliminar_val,
                'imprimir': imprimir_val,
            }

        perfil.permisos_detalle = detalle
        perfil.permiso_dashboard = detalle['dashboard']['activo']
        perfil.permiso_inventario = detalle['inventario']['activo']
        perfil.permiso_entregas = detalle['entregas']['activo']
        perfil.permiso_atencion_ciudadano = detalle['atencion_ciudadano']['activo']
        perfil.permiso_asignaciones_especiales = detalle['asignaciones_especiales']['activo']
        perfil.permiso_usuarios = detalle['usuarios']['activo']
        perfil.permiso_auditoria = detalle['auditoria']['activo']
        perfil.save()

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Permisos manuales configurados para '{u.username}' (Rol: {rol_nombre}) con {perfil.total_modulos_activos}/10 módulos activos.",
            modulo="Usuarios",
            timestamp=timezone.now()
        )

        messages.success(request, f"🛡️ Permisos manuales de '{u.username}' actualizados exitosamente ({perfil.total_modulos_activos}/10 módulos habilitados).")
    return redirect('registrar_usuario')

@login_required
def cambiar_rol_usuario(request, user_id):
    if not request.user.is_superuser:
        messages.error(request, "Acceso restringido: Solo los Administradores pueden modificar roles.")
        return redirect('dashboard')

    if request.method == 'POST':
        u = get_object_or_404(User, id=user_id)
        nuevo_rol = request.POST.get('nuevo_rol', 'inventario')

        if u.id == request.user.id and nuevo_rol != 'admin':
            messages.error(request, "No puedes quitarte el rol de Administrador a ti mismo.")
            return redirect('registrar_usuario')

        if nuevo_rol == 'admin':
            u.is_superuser = True
            u.is_staff = True
        elif nuevo_rol in ['inventario', 'entregas']:
            u.is_superuser = False
            u.is_staff = True
        else:
            u.is_superuser = False
            u.is_staff = False

        u.first_name = nuevo_rol
        u.save()

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Cambio de rol para usuario '{u.username}' a '{nuevo_rol.upper()}'.",
            modulo="Usuarios",
            timestamp=timezone.now()
        )

        messages.success(request, f"Rol de '{u.username}' actualizado a {nuevo_rol.upper()}.")
    return redirect('registrar_usuario')

@login_required
def alternar_estado_usuario(request, user_id):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('usuarios'))):
        messages.error(request, "Acceso restringido: Solo los Administradores o usuarios autorizados pueden desactivar cuentas.")
        return redirect('dashboard')

    if request.method == 'POST':
        u = get_object_or_404(User, id=user_id)

        if u.id == request.user.id:
            messages.error(request, "No puedes desactivar tu propia cuenta.")
            return redirect('registrar_usuario')

        u.is_active = not u.is_active
        u.save()

        estado_txt = "ACTIVADA" if u.is_active else "DESACTIVADA"
        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Cuenta de usuario '{u.username}' {estado_txt}.",
            modulo="Usuarios",
            timestamp=timezone.now()
        )

        messages.success(request, f"Cuenta de '{u.username}' {estado_txt}.")
    return redirect('registrar_usuario')

@login_required
def vista_usuarios_conectados(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('usuarios'))):
        messages.error(request, "Acceso restringido: Solo los Administradores o usuarios autorizados pueden ver usuarios conectados.")
        return redirect('dashboard')
    return render(request, 'usuarios_conectados.html')

@login_required
def api_usuarios_conectados(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('usuarios'))):
        return JsonResponse({'error': 'No autorizado'}, status=403)
    usuarios = User.objects.filter(is_active=True)
    conectados = []
    
    for u in usuarios:
        cache_key = f'last_seen_{u.id}'
        last_seen = cache.get(cache_key)
        
        # Consideramos conectado si tuvo actividad en los últimos 15 minutos (900 seg)
        if last_seen and (timezone.now() - last_seen).total_seconds() < 900:
            rol_nombre = u.perfil.rol if hasattr(u, 'perfil') else ('Administrador' if u.is_superuser else 'Usuario')
            conectados.append({
                'id': u.id,
                'username': u.username,
                'nombre_completo': u.get_full_name() or u.username,
                'rol': rol_nombre,
                'is_superuser': u.is_superuser,
                'ultima_actividad': last_seen.strftime('%d/%m/%Y %H:%M:%S'),
                'hace_minutos': int((timezone.now() - last_seen).total_seconds() // 60)
            })
            
    return JsonResponse({'usuarios': conectados})

@login_required
def lista_auditoria(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('auditoria'))):
        messages.error(request, "Acceso denegado: Únicamente los administradores o usuarios autorizados pueden consultar el registro de auditoría.")
        return redirect('dashboard')

    modulo_filtro = request.GET.get('modulo', 'Todos').strip()
    q = request.GET.get('q', '').strip()

    logs = RegistroAuditoria.objects.all().order_by('-timestamp')

    if modulo_filtro and modulo_filtro != 'Todos':
        logs = logs.filter(modulo=modulo_filtro)

    if q:
        logs = logs.filter(
            Q(accion__icontains=q) |
            Q(usuario__username__icontains=q) |
            Q(modulo__icontains=q)
        )

    per_page_raw = request.GET.get('per_page', '50').strip()
    try:
        per_page = int(per_page_raw)
        if per_page not in [15, 30, 50, 100]:
            per_page = 50
    except ValueError:
        per_page = 50

    paginator = Paginator(logs, per_page)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    return render(request, 'lista_auditoria.html', {
        'registros': page_obj.object_list,
        'logs': page_obj.object_list,
        'page_obj': page_obj,
        'modulo_actual': modulo_filtro,
        'per_page': per_page,
        'q': q
    })

@login_required
def exportar_auditoria_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log Auditoria"

    ws.merge_cells('A1:D1')
    ws['A1'] = "SIA WEB - REPORTE DE AUDITORÍA Y SEGURIDAD"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = ["Fecha y Hora", "Usuario", "Módulo", "Acción Realizada"]
    ws.append(headers)
    ws.row_dimensions[2].height = 24

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    logs = RegistroAuditoria.objects.all().order_by('-timestamp')
    for row_idx, l in enumerate(logs, start=3):
        fecha_str = l.timestamp.strftime('%d/%m/%Y %H:%M:%S') if l.timestamp else 'N/A'
        user_str = l.usuario.username if l.usuario else 'Sistema'
        ws.append([fecha_str, user_str, l.modulo, l.accion])
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Reporte_Auditoria.xlsx"'
    wb.save(response)
    return response

# --- MÓDULO DE ATENCIÓN AL CIUDADANO ---

@login_required
def lista_solicitudes_ciudadano(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_atencion_ciudadano:
            messages.error(request, "Acceso restringido: No tienes permiso para acceder al módulo de Atención al Ciudadano.")
            return redirect(obtener_url_inicio_usuario(request.user))

    status_filtro = request.GET.get('status', 'Todos').strip()
    prioridad_filtro = request.GET.get('prioridad', 'Todas').strip()
    fecha_inicio_raw = request.GET.get('fecha_inicio', '').strip()
    fecha_fin_raw = request.GET.get('fecha_fin', '').strip()
    q = request.GET.get('q', '').strip()

    qs = SolicitudCiudadano.objects.all()

    if status_filtro and status_filtro != 'Todos':
        qs = qs.filter(status=status_filtro)

    if prioridad_filtro and prioridad_filtro != 'Todas':
        qs = qs.filter(prioridad=prioridad_filtro)

    fecha_inicio_val = parse_date_safe(fecha_inicio_raw)
    if fecha_inicio_val:
        qs = qs.filter(fecha_solicitud__date__gte=fecha_inicio_val)

    fecha_fin_val = parse_date_safe(fecha_fin_raw)
    if fecha_fin_val:
        qs = qs.filter(fecha_solicitud__date__lte=fecha_fin_val)

    if q:
        qs = qs.filter(
            Q(cedula__icontains=q) |
            Q(nombre_apellido__icontains=q) |
            Q(tipo_solicitud__icontains=q) |
            Q(descripcion_solicitud__icontains=q)
        )

    solicitudes = qs.order_by('-fecha_solicitud')

    total_solicitudes = SolicitudCiudadano.objects.count()
    recibidas_count = SolicitudCiudadano.objects.filter(status='Recibida').count()
    en_proceso_count = SolicitudCiudadano.objects.filter(status='En Proceso').count()
    aprobadas_count = SolicitudCiudadano.objects.filter(status='Aprobada').count()

    per_page_raw = request.GET.get('per_page', '30').strip()
    try:
        per_page = int(per_page_raw)
        if per_page not in [15, 30, 50, 100]:
            per_page = 30
    except ValueError:
        per_page = 30

    paginator = Paginator(solicitudes, per_page)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    lista_inventario = list(SIA_producto.objects.filter(cantidad__gt=0).values('id', 'codigo', 'descripcion', 'cantidad', 'almacen').order_by('descripcion'))
    hoy = timezone.now().date().strftime('%Y-%m-%d')

    context = {
        'solicitudes': page_obj.object_list,
        'page_obj': page_obj,
        'status_actual': status_filtro,
        'prioridad_actual': prioridad_filtro,
        'fecha_inicio': fecha_inicio_raw,
        'fecha_fin': fecha_fin_raw,
        'per_page': per_page,
        'total_solicitudes': total_solicitudes,
        'recibidas_count': recibidas_count,
        'en_proceso_count': en_proceso_count,
        'aprobadas_count': aprobadas_count,
        'lista_inventario': lista_inventario,
        'hoy': hoy,
        'q': q
    }
    return render(request, 'atencion_ciudadano/lista_solicitudes.html', context)

@login_required
def exportar_solicitudes_excel(request):
    status_filtro = request.GET.get('status', 'Todos').strip()
    prioridad_filtro = request.GET.get('prioridad', 'Todas').strip()
    fecha_inicio_raw = request.GET.get('fecha_inicio', '').strip()
    fecha_fin_raw = request.GET.get('fecha_fin', '').strip()
    q = request.GET.get('q', '').strip()

    qs = SolicitudCiudadano.objects.all()

    if status_filtro and status_filtro != 'Todos':
        qs = qs.filter(status=status_filtro)

    if prioridad_filtro and prioridad_filtro != 'Todas':
        qs = qs.filter(prioridad=prioridad_filtro)

    fecha_inicio_val = parse_date_safe(fecha_inicio_raw)
    if fecha_inicio_val:
        qs = qs.filter(fecha_solicitud__date__gte=fecha_inicio_val)

    fecha_fin_val = parse_date_safe(fecha_fin_raw)
    if fecha_fin_val:
        qs = qs.filter(fecha_solicitud__date__lte=fecha_fin_val)

    if q:
        qs = qs.filter(
            Q(cedula__icontains=q) |
            Q(nombre_apellido__icontains=q) |
            Q(tipo_solicitud__icontains=q) |
            Q(descripcion_solicitud__icontains=q)
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solicitudes Ciudadano"

    ws.merge_cells('A1:J1')
    ws['A1'] = "SIA WEB - REPORTE DE ATENCIÓN AL CIUDADANO"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = ["# Solicitud", "Fecha", "Cédula", "Solicitante", "Teléfono", "Categoría", "Requerimiento", "Prioridad", "Estado", "Observaciones"]
    ws.append(headers)
    ws.row_dimensions[2].height = 24

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    solicitudes = qs.order_by('-fecha_solicitud')
    for row_idx, s in enumerate(solicitudes, start=3):
        fecha_str = s.fecha_solicitud.strftime('%d/%m/%Y %H:%M') if s.fecha_solicitud else 'N/A'
        ws.append([
            f"#{s.id}",
            fecha_str,
            s.cedula,
            s.nombre_apellido,
            s.telefono or 'N/A',
            s.tipo_solicitud,
            s.descripcion_solicitud,
            s.prioridad,
            s.status,
            s.observaciones_seguimiento or ''
        ])
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=8).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=9).alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Reporte_Atencion_Ciudadano.xlsx"'
    wb.save(response)
    return response

@login_required
def generar_reporte_solicitudes_pdf(request):
    status_filtro = request.GET.get('status', 'Todos').strip()
    prioridad_filtro = request.GET.get('prioridad', 'Todas').strip()
    fecha_inicio_raw = request.GET.get('fecha_inicio', '').strip()
    fecha_fin_raw = request.GET.get('fecha_fin', '').strip()
    q = request.GET.get('q', '').strip()

    qs = SolicitudCiudadano.objects.all()

    if status_filtro and status_filtro != 'Todos':
        qs = qs.filter(status=status_filtro)

    if prioridad_filtro and prioridad_filtro != 'Todas':
        qs = qs.filter(prioridad=prioridad_filtro)

    fecha_inicio_val = parse_date_safe(fecha_inicio_raw)
    if fecha_inicio_val:
        qs = qs.filter(fecha_solicitud__date__gte=fecha_inicio_val)

    fecha_fin_val = parse_date_safe(fecha_fin_raw)
    if fecha_fin_val:
        qs = qs.filter(fecha_solicitud__date__lte=fecha_fin_val)

    if q:
        qs = qs.filter(
            Q(cedula__icontains=q) |
            Q(nombre_apellido__icontains=q) |
            Q(tipo_solicitud__icontains=q) |
            Q(descripcion_solicitud__icontains=q)
        )

    solicitudes = qs.order_by('-fecha_solicitud')
    total_count = solicitudes.count()
    recibidas_count = solicitudes.filter(status='Recibida').count()
    en_proceso_count = solicitudes.filter(status='En Proceso').count()
    aprobadas_count = solicitudes.filter(status='Aprobada').count()
    entregadas_count = solicitudes.filter(status='Entregada').count()

    from django.template.loader import render_to_string
    from xhtml2pdf import pisa

    context = {
        'solicitudes': solicitudes,
        'total_count': total_count,
        'recibidas_count': recibidas_count,
        'en_proceso_count': en_proceso_count,
        'aprobadas_count': aprobadas_count,
        'entregadas_count': entregadas_count,
        'status_filtro': status_filtro,
        'prioridad_filtro': prioridad_filtro,
        'fecha_inicio': fecha_inicio_raw,
        'fecha_fin': fecha_fin_raw,
        'fecha_generacion': timezone.now().strftime('%d/%m/%Y %H:%M'),
        'usuario_generador': request.user.username
    }

    html = render_to_string('reportes/reporte_solicitudes_pdf.html', context)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="Reporte_Atencion_Ciudadano.pdf"'
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
        return HttpResponse('Error al generar el PDF de solicitudes', status=500)
    return response

@login_required
def registrar_solicitud_ciudadano(request):
    if request.method == 'POST':
        cedula = request.POST.get('cedula', '').strip()
        nombre_apellido = request.POST.get('nombre_apellido', '').strip()
        telefono = request.POST.get('telefono', '').strip()
        direccion = request.POST.get('direccion', '').strip()
        tipo_solicitud = request.POST.get('tipo_solicitud', '').strip()
        descripcion_solicitud = request.POST.get('descripcion_solicitud', '').strip()
        prioridad = request.POST.get('prioridad', 'Media').strip()

        if not cedula or not nombre_apellido or not tipo_solicitud or not descripcion_solicitud:
            messages.error(request, "Por favor completa todos los campos obligatorios del formulario (Cédula, Nombre, Categoría y Requerimiento).")
            return redirect('registrar_solicitud_ciudadano')

        # Actualizar o crear ficha de beneficiario/ciudadano
        Beneficiario.objects.update_or_create(
            cedula=cedula,
            defaults={
                'nombre_apellido': nombre_apellido,
                'telefono': telefono,
                'direccion': direccion
            }
        )

        solicitud = SolicitudCiudadano.objects.create(
            cedula=cedula,
            nombre_apellido=nombre_apellido,
            telefono=telefono,
            direccion=direccion,
            tipo_solicitud=tipo_solicitud,
            descripcion_solicitud=descripcion_solicitud,
            prioridad=prioridad,
            registrado_por=request.user
        )

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Registro de solicitud de atención al ciudadano #{solicitud.id} para {nombre_apellido} (C.I.: {cedula}) - {tipo_solicitud}.",
            modulo="Atención al Ciudadano",
            timestamp=timezone.now()
        )

        messages.success(request, f"Solicitud #{solicitud.id} registrada correctamente para {nombre_apellido}.")
        return redirect('lista_solicitudes_ciudadano')

    return render(request, 'atencion_ciudadano/registrar_solicitud.html')

@login_required
def cambiar_estado_solicitud(request, pk):
    if not request.user.is_superuser:
        messages.error(request, "Acceso denegado: Únicamente los administradores pueden gestionar el estado y seguimiento de las solicitudes.")
        return redirect('lista_solicitudes_ciudadano')

    solicitud = get_object_or_404(SolicitudCiudadano, id=pk)
    if solicitud.status in ['Entregada', 'Entregado']:
        messages.warning(request, f"La solicitud #{solicitud.id} ya se encuentra en estado 'Entregada' y no puede ser modificada.")
        return redirect('lista_solicitudes_ciudadano')

    if request.method == 'POST':
        nuevo_status = request.POST.get('status', solicitud.status)
        nueva_prioridad = request.POST.get('prioridad', solicitud.prioridad)
        observaciones = request.POST.get('observaciones_seguimiento', '').strip()

        solicitud.status = nuevo_status
        solicitud.prioridad = nueva_prioridad
        if observaciones:
            solicitud.observaciones_seguimiento = observaciones
        solicitud.save()

        # Enviar notificación por correo si aplica
        enviar_notificacion_solicitud_email(solicitud, nuevo_status)

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Actualización de estado en solicitud #{solicitud.id} a '{nuevo_status}' (Prioridad: {nueva_prioridad}).",
            modulo="Atención al Ciudadano",
            timestamp=timezone.now()
        )

        messages.success(request, f"Estado de la solicitud #{solicitud.id} actualizado correctamente a '{nuevo_status}'.")
    return redirect('lista_solicitudes_ciudadano')

@login_required
def entregar_solicitud_desde_inventario(request, pk):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_atencion_ciudadano:
            messages.error(request, "Acceso restringido: No tienes permiso para gestionar entregas en Atención al Ciudadano.")
            return redirect('lista_solicitudes_ciudadano')

    solicitud = get_object_or_404(SolicitudCiudadano, id=pk)

    if solicitud.status in ['Entregada', 'Entregado']:
        messages.warning(request, f"La solicitud #{solicitud.id} ya figura como entregada.")
        return redirect('lista_solicitudes_ciudadano')

    if request.method == 'POST':
        producto_id = request.POST.get('producto_id')
        cantidad_str = request.POST.get('cantidad', '1')
        fecha_entrega_str = request.POST.get('fecha_entrega')
        observaciones = request.POST.get('observaciones_seguimiento', '').strip()

        try:
            cantidad = int(cantidad_str)
            if cantidad <= 0:
                raise ValueError()
        except ValueError:
            messages.error(request, "La cantidad a entregar debe ser un número entero positivo mayor a 0.")
            return redirect('lista_solicitudes_ciudadano')

        fecha_entrega = parse_date_safe(fecha_entrega_str) or timezone.now().date()

        # Subida segura de evidencias fotográficas (hasta 3) FUERA de la transacción de base de datos
        archivos = [
            request.FILES.get('evidencia1'),
            request.FILES.get('evidencia2'),
            request.FILES.get('evidencia3'),
        ]
        urls_evidencia = []
        for arc in archivos:
            if arc:
                url_subida = subir_archivo_evidencia_seguro(arc, request)
                urls_evidencia.append(url_subida)
            else:
                urls_evidencia.append(None)

        with transaction.atomic():
            prod = SIA_producto.objects.select_for_update().filter(id=producto_id).first()
            if not prod:
                messages.error(request, "El producto seleccionado del inventario no existe o no fue especificado.")
                return redirect('lista_solicitudes_ciudadano')

            if prod.cantidad < cantidad:
                messages.error(
                    request,
                    f"Stock insuficiente para '{prod.descripcion}'. Solicitado: {cantidad} u., Disponible en inventario: {prod.cantidad} u."
                )
                return redirect('lista_solicitudes_ciudadano')

            # Descontar stock atómicamente
            prod.cantidad -= cantidad
            prod.save(update_fields=['cantidad'])

            # Sincronizar directorio de beneficiarios
            cedula_clean = solicitud.cedula.strip() if solicitud.cedula else ''
            if cedula_clean:
                beneficiario_obj = Beneficiario.objects.filter(cedula=cedula_clean).first()
                if not beneficiario_obj and solicitud.nombre_apellido:
                    Beneficiario.objects.create(
                        cedula=cedula_clean,
                        nombre_apellido=solicitud.nombre_apellido,
                        direccion=solicitud.direccion or '',
                        telefono=solicitud.telefono or '',
                        timestamp=timezone.now()
                    )

            # Crear entrega en BeneficioEntregado
            entrega = BeneficioEntregado.objects.create(
                producto=prod,
                cantidad_dada=cantidad,
                cedula=cedula_clean,
                codigo_prod=prod.codigo or '',
                descripcion_prod=prod.descripcion or '',
                fecha=timezone.now().date(),
                fecha_entrega=fecha_entrega,
                nombre_beneficiario=solicitud.nombre_apellido,
                status='Entregado',
                timestamp=timezone.now(),
                url_evidencia_1=urls_evidencia[0],
                url_evidencia_2=urls_evidencia[1],
                url_evidencia_3=urls_evidencia[2],
                via=f"Solicitud #{solicitud.id} - Atención al Ciudadano"
            )

            # Actualizar SolicitudCiudadano
            solicitud.entrega = entrega
            solicitud.status = 'Entregada'
            solicitud.fecha_entrega = fecha_entrega
            solicitud.url_evidencia_1 = urls_evidencia[0]
            solicitud.url_evidencia_2 = urls_evidencia[1]
            solicitud.url_evidencia_3 = urls_evidencia[2]
            if observaciones:
                if solicitud.observaciones_seguimiento:
                    solicitud.observaciones_seguimiento += f"\n[Entrega]: {observaciones}"
                else:
                    solicitud.observaciones_seguimiento = f"[Entrega]: {observaciones}"
            solicitud.save()

            # Registro de auditoría
            RegistroAuditoria.objects.create(
                usuario=request.user,
                accion=f"Entrega directa desde Inventario de solicitud #{solicitud.id}: {cantidad}x {prod.descripcion} a {solicitud.nombre_apellido} (C.I. {solicitud.cedula}). Stock restante: {prod.cantidad} u.",
                modulo="Atención al Ciudadano",
                timestamp=timezone.now()
            )

            # Notificación por correo si aplica
            enviar_notificacion_solicitud_email(solicitud, 'Entregada')

            messages.success(
                request,
                f"✅ Solicitud #{solicitud.id} entregada con éxito: {cantidad} u. de '{prod.descripcion}' descontadas del inventario."
            )

    return redirect('lista_solicitudes_ciudadano')

@login_required
def eliminar_solicitud_ciudadano(request, pk):
    if not request.user.is_superuser:
        messages.error(request, "Únicamente administradores pueden eliminar registros de atención al ciudadano.")
        return redirect('lista_solicitudes_ciudadano')

    solicitud = get_object_or_404(SolicitudCiudadano, id=pk)
    if request.method == 'POST':
        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Eliminación de solicitud de atención al ciudadano #{solicitud.id} ({solicitud.nombre_apellido}).",
            modulo="Atención al Ciudadano",
            timestamp=timezone.now()
        )
        solicitud.delete()
        messages.success(request, "Solicitud eliminada exitosamente.")
    return redirect('lista_solicitudes_ciudadano')

def obtener_ip_red_local():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

def obtener_url_publica_comprobante(request, custom_path=None):
    target_path = custom_path if custom_path else request.path
    host = request.get_host()
    if 'trycloudflare.com' in host or 'render.com' in host or 'herokuapp.com' in host:
        scheme = 'https' if request.is_secure() or request.headers.get('x-forwarded-proto') == 'https' else 'http'
        return f"{scheme}://{host}{target_path}"

    tunnel_file = os.path.join(settings.BASE_DIR, '.active_tunnel')
    if os.path.isfile(tunnel_file):
        try:
            mtime = os.path.getmtime(tunnel_file)
            if (timezone.now().timestamp() - mtime) < 28800:
                with open(tunnel_file, 'r', encoding='utf-8') as f:
                    tunnel_url = f.read().strip()
                if tunnel_url.startswith('http'):
                    return f"{tunnel_url.rstrip('/')}{target_path}"
        except Exception:
            pass

    url_base = request.build_absolute_uri(target_path)
    if '127.0.0.1' in host or 'localhost' in host:
        lan_ip = obtener_ip_red_local()
        puerto = host.split(':')[1] if ':' in host else '8000'
        url_base = url_base.replace(f'127.0.0.1:{puerto}', f"{lan_ip}:{puerto}").replace(f'localhost:{puerto}', f"{lan_ip}:{puerto}")
        url_base = url_base.replace('127.0.0.1', lan_ip).replace('localhost', lan_ip)

    return url_base

def imprimir_comprobante_solicitud(request, pk):
    solicitud = get_object_or_404(SolicitudCiudadano, id=pk)
    url_base = obtener_url_publica_comprobante(request)
    firma_hash = generar_hash_firma_digital(solicitud)

    return render(request, 'atencion_ciudadano/comprobante_solicitud.html', {
        'solicitud': solicitud,
        'qr_comprobante_url': url_base,
        'firma_hash': firma_hash
    })

# Módulo de Respaldo y Restauración de Base de Datos (Administradores / Delegados)
@login_required
def gestion_respaldos(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('respaldos'))):
        messages.error(request, "Acceso restringido: No tienes permiso para gestionar los respaldos del sistema.")
        return redirect(obtener_url_inicio_usuario(request.user))
        
    from core.utils import ejecutar_respaldo_automatico_si_aplica
    ejecutar_respaldo_automatico_si_aplica(request)

    backup_dir = os.path.join(settings.BASE_DIR, 'backups')
    os.makedirs(backup_dir, exist_ok=True)
    
    archivos = []
    total_bytes = 0
    
    formatos_permitidos = ('.json', '.sqlite3', '.sql', '.dump', '.gz')
    for fname in os.listdir(backup_dir):
        if any(fname.endswith(ext) for ext in formatos_permitidos):
            fpath = os.path.join(backup_dir, fname)
            if os.path.isfile(fpath):
                stat = os.stat(fpath)
                size_mb = round(stat.st_size / (1024 * 1024), 2)
                total_bytes += stat.st_size
                mtime = datetime.fromtimestamp(stat.st_mtime)
                
                archivos.append({
                    'nombre': fname,
                    'tamano_mb': size_mb,
                    'fecha': mtime,
                    'es_json': fname.endswith('.json'),
                    'es_sqlite': fname.endswith('.sqlite3'),
                    'es_pg': fname.endswith('.sql') or fname.endswith('.dump') or fname.endswith('.gz')
                })
                
    archivos.sort(key=lambda x: x['fecha'], reverse=True)
    total_mb = round(total_bytes / (1024 * 1024), 2)
    ultimo_respaldo = archivos[0]['fecha'] if archivos else None
    
    return render(request, 'admin_sistema/gestion_respaldos.html', {
        'archivos': archivos,
        'total_respaldos': len(archivos),
        'total_mb': total_mb,
        'ultimo_respaldo': ultimo_respaldo
    })

@login_required
def generar_respaldo_manual_view(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('respaldos', 'crear'))):
        messages.error(request, "Acceso restringido: No tienes permiso para generar respaldos.")
        return redirect('gestion_respaldos')
        
    try:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        backup_dir = os.path.join(settings.BASE_DIR, 'backups')
        os.makedirs(backup_dir, exist_ok=True)

        archivos_generados = []

        # 1. Respaldo universal JSON mediante dumpdata
        json_file = os.path.join(backup_dir, f'respaldo_sia_{timestamp}.json')
        with open(json_file, 'w', encoding='utf-8') as f:
            call_command('dumpdata', '--natural-foreign', '--natural-primary', exclude=['contenttypes', 'auth.permission'], stdout=f)
        archivos_generados.append(os.path.basename(json_file))

        # 2. Respaldo nativo PostgreSQL con pg_dump
        db_conf = settings.DATABASES.get('default', {})
        if 'postgresql' in db_conf.get('ENGINE', ''):
            pg_file = os.path.join(backup_dir, f'pg_dump_sia_{timestamp}.dump')
            env = os.environ.copy()
            if db_conf.get('PASSWORD'):
                env['PGPASSWORD'] = str(db_conf['PASSWORD'])
            cmd = [
                'pg_dump',
                '-h', str(db_conf.get('HOST') or '127.0.0.1'),
                '-p', str(db_conf.get('PORT') or '5432'),
                '-U', str(db_conf.get('USER') or 'postgres'),
                '-d', str(db_conf.get('NAME') or 'inventario'),
                '-F', 'c',
                '-f', pg_file
            ]
            try:
                import subprocess
                res = subprocess.run(cmd, env=env, capture_output=True, timeout=25)
                if res.returncode == 0 and os.path.exists(pg_file) and os.path.getsize(pg_file) > 0:
                    archivos_generados.append(os.path.basename(pg_file))
            except Exception:
                pass

        # 3. Respaldo SQLite si la base de datos es SQLite
        sqlite_file = os.path.join(settings.BASE_DIR, 'db.sqlite3')
        if os.path.exists(sqlite_file) and 'sqlite' in db_conf.get('ENGINE', ''):
            sqlite_copy = os.path.join(backup_dir, f'db_copy_{timestamp}.sqlite3')
            shutil.copy2(sqlite_file, sqlite_copy)
            archivos_generados.append(os.path.basename(sqlite_copy))

        archivos_txt = ", ".join(archivos_generados)
        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion=f"Generación manual de respaldo de base de datos ({archivos_txt}).",
            modulo="Administración",
            timestamp=timezone.now()
        )
        messages.success(request, f"Copia de seguridad generada con éxito: {archivos_txt}.")
    except Exception as e:
        messages.error(request, f"Error al generar respaldo: {str(e)}")
        
    return redirect('gestion_respaldos')

@login_required
def descargar_respaldo(request, filename):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('respaldos', 'leer'))):
        messages.error(request, "Acceso restringido.")
        return redirect('gestion_respaldos')
        
    filename = os.path.basename(filename)
    fpath = os.path.join(settings.BASE_DIR, 'backups', filename)
    
    if not os.path.exists(fpath):
        raise Http404("El archivo de respaldo no existe.")
        
    return FileResponse(open(fpath, 'rb'), as_attachment=True, filename=filename)

@login_required
def restaurar_respaldo_view(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('respaldos', 'crear'))):
        messages.error(request, "Acceso restringido: Solo administradores autorizados pueden restaurar respaldos.")
        return redirect('gestion_respaldos')
        
    if request.method == 'POST':
        file_to_restore = None
        
        # 1. Si el usuario subió un archivo desde su equipo
        if 'archivo_respaldo' in request.FILES:
            uploaded_file = request.FILES['archivo_respaldo']
            if not (uploaded_file.name.endswith('.json') or uploaded_file.name.endswith('.sqlite3')):
                messages.error(request, "Formato no válido. Solo se permiten archivos .json o .sqlite3.")
                return redirect('gestion_respaldos')
                
            backup_dir = os.path.join(settings.BASE_DIR, 'backups')
            os.makedirs(backup_dir, exist_ok=True)
            save_path = os.path.join(backup_dir, f"subido_{uploaded_file.name}")
            with open(save_path, 'wb+') as destination:
                for chunk in uploaded_file.chunks():
                    destination.write(chunk)
            file_to_restore = save_path
        
        # 2. O si seleccionó un archivo ya existente en servidor
        elif 'nombre_archivo' in request.POST:
            fname = os.path.basename(request.POST.get('nombre_archivo', '').strip())
            file_to_restore = os.path.join(settings.BASE_DIR, 'backups', fname)
            
        if not file_to_restore or not os.path.exists(file_to_restore):
            messages.error(request, "No se encontró el archivo de respaldo especificado para la restauración.")
            return redirect('gestion_respaldos')
            
        try:
            if file_to_restore.endswith('.json'):
                call_command('loaddata', file_to_restore)
                msg = f"Base de datos restaurada exitosamente desde el archivo JSON '{os.path.basename(file_to_restore)}'."
            elif file_to_restore.endswith('.sqlite3'):
                target_sqlite = os.path.join(settings.BASE_DIR, 'db.sqlite3')
                shutil.copy2(file_to_restore, target_sqlite)
                msg = f"Base de datos SQLite reemplazada exitosamente desde la copia '{os.path.basename(file_to_restore)}'."
            else:
                messages.error(request, "Formato de archivo no soportado para la restauración.")
                return redirect('gestion_respaldos')

            RegistroAuditoria.objects.create(
                usuario=request.user,
                accion=f"Restauración de base de datos ejecutada con éxito usando '{os.path.basename(file_to_restore)}'.",
                modulo="Administración",
                timestamp=timezone.now()
            )
            messages.success(request, msg)
        except Exception as e:
            messages.error(request, f"Error al restaurar la base de datos: {str(e)}")
            
    return redirect('gestion_respaldos')

@login_required
def eliminar_respaldo_view(request, filename):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and request.user.perfil.tiene_permiso('respaldos', 'eliminar'))):
        messages.error(request, "Acceso restringido: No tienes permiso para eliminar respaldos de la base de datos.")
        return redirect('gestion_respaldos')
        
    fname = os.path.basename(filename)
    fpath = os.path.join(settings.BASE_DIR, 'backups', fname)
    
    if os.path.exists(fpath) and os.path.isfile(fpath):
        try:
            os.remove(fpath)
            RegistroAuditoria.objects.create(
                usuario=request.user,
                accion=f"Eliminación de copia de respaldo de base de datos ({fname}).",
                modulo="Administración",
                timestamp=timezone.now()
            )
            messages.success(request, f"🗑️ El respaldo '{fname}' ha sido eliminado exitosamente.")
        except Exception as e:
            messages.error(request, f"Error al eliminar el archivo de respaldo: {str(e)}")
    else:
        messages.error(request, f"El archivo de respaldo '{fname}' no existe o ya fue eliminado.")
        
    return redirect('gestion_respaldos')

# --- MÓDULO DE FICHA 360 Y ESTADÍSTICAS AVANZADAS ---
@login_required
def ficha_ciudadano_360(request, cedula):
    cedula_clean = str(cedula).strip()
    beneficiario = Beneficiario.objects.filter(cedula=cedula_clean).first()
    
    solicitudes = SolicitudCiudadano.objects.filter(cedula=cedula_clean).order_by('-fecha_solicitud')
    entregas = BeneficioEntregado.objects.filter(cedula=cedula_clean).order_by('-id')
    
    total_solicitudes = solicitudes.count()
    total_entregas = entregas.count()
    
    return render(request, 'atencion_ciudadano/ficha_360.html', {
        'cedula': cedula_clean,
        'beneficiario': beneficiario,
        'solicitudes': solicitudes,
        'entregas': entregas,
        'total_solicitudes': total_solicitudes,
        'total_entregas': total_entregas
    })

@login_required
def panel_estadisticas_avanzadas(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.permiso_dashboard:
            messages.error(request, "Acceso restringido: No tienes permiso para ver el panel de estadísticas.")
            return redirect(obtener_url_inicio_usuario(request.user))
            
    por_categoria = SolicitudCiudadano.objects.values('tipo_solicitud').annotate(total=Count('id')).order_by('-total')
    cat_labels = [c['tipo_solicitud'] for c in por_categoria if c['tipo_solicitud']]
    cat_totals = [c['total'] for c in por_categoria if c['tipo_solicitud']]
    
    por_estado = SolicitudCiudadano.objects.values('status').annotate(total=Count('id'))
    estado_dict = {e['status']: e['total'] for e in por_estado if e['status']}
    
    cnt_en_evaluacion = estado_dict.get('En Proceso', 0) + estado_dict.get('Recibida', 0) + estado_dict.get('Aprobada', 0)
    
    top_productos = BeneficioEntregado.objects.values('descripcion_prod').annotate(total_cant=Sum('cantidad_dada')).order_by('-total_cant')[:6]
    prod_labels = [p['descripcion_prod'] or 'Otros' for p in top_productos]
    prod_totals = [p['total_cant'] for p in top_productos]
    
    por_comunidad = SolicitudCiudadano.objects.values('direccion').annotate(total=Count('id')).order_by('-total')[:6]
    comunidad_labels = [c['direccion'] or 'No especificado' for c in por_comunidad]
    comunidad_totals = [c['total'] for c in por_comunidad]

    total_solicitudes = SolicitudCiudadano.objects.count()
    total_entregas = BeneficioEntregado.objects.count()
    
    meses_nombres = {1: 'Ene', 2: 'Feb', 3: 'Mar', 4: 'Abr', 5: 'May', 6: 'Jun', 7: 'Jul', 8: 'Ago', 9: 'Sep', 10: 'Oct', 11: 'Nov', 12: 'Dic'}
    now = timezone.now()
    meses_labels = []
    solicitudes_mensuales = []
    entregas_mensuales = []
    
    for i in range(5, -1, -1):
        year = now.year
        month = now.month - i
        if month <= 0:
            month += 12
            year -= 1
            
        meses_labels.append(f"{meses_nombres[month]} {year}")
        solicitudes_mensuales.append(SolicitudCiudadano.objects.filter(fecha_solicitud__year=year, fecha_solicitud__month=month).count())
        entregas_mensuales.append(BeneficioEntregado.objects.filter(fecha__year=year, fecha__month=month).count())

    eficiencia_pct = round((total_entregas / total_solicitudes * 100), 1) if total_solicitudes > 0 else 100.0

    return render(request, 'reportes/panel_estadisticas.html', {
        'cat_labels_json': json.dumps(cat_labels),
        'cat_totals_json': json.dumps(cat_totals),
        'cnt_en_evaluacion': cnt_en_evaluacion,
        'prod_labels_json': json.dumps(prod_labels),
        'prod_totals_json': json.dumps(prod_totals),
        'comunidad_labels_json': json.dumps(comunidad_labels),
        'comunidad_totals_json': json.dumps(comunidad_totals),
        'meses_labels_json': json.dumps(meses_labels),
        'solicitudes_mensuales_json': json.dumps(solicitudes_mensuales),
        'entregas_mensuales_json': json.dumps(entregas_mensuales),
        'eficiencia_pct': eficiencia_pct,
        'total_solicitudes': total_solicitudes,
        'total_entregas': total_entregas
    })

@login_required
def generar_reporte_parroquias_pdf(request):
    if not request.user.is_superuser:
        perfil = getattr(request.user, 'perfil', None)
        if perfil and not perfil.puede_imprimir_estadisticas:
            messages.error(request, "Acceso restringido: No tienes permiso para descargar reportes PDF.")
            return redirect('dashboard')

    comunidades = SolicitudCiudadano.objects.values('direccion').annotate(
        total_solicitudes=Count('id'),
        entregadas=Count('id', filter=Q(status__iexact='Entregada')),
        en_proceso=Count('id', filter=Q(status__iexact='En Proceso'))
    ).order_by('-total_solicitudes')

    html_string = render_to_string('reportes/reporte_parroquias_pdf.html', {
        'comunidades': comunidades,
        'fecha_reporte': datetime.now(),
        'usuario_impresion': request.user.username,
        'firma_digital': generar_hash_firma_digital(request.user)
    })
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_sectores_parroquias.pdf"'
    pisa_status = pisa.CreatePDF(html_string, dest=response)
    
    if pisa_status.err:
        return HttpResponse('Error al generar PDF de sectores', status=500)
    return response

@login_required
def generar_constancia_solicitud_pdf(request, solicitud_id):
    solicitud = get_object_or_404(SolicitudCiudadano, id=solicitud_id)
    
    server_ip = request.META.get('HTTP_HOST', 'localhost:8000')
    scheme = 'https' if request.is_secure() else 'http'
    qr_url = f"{scheme}://{server_ip}/atencion-ciudadano/{solicitud.id}/imprimir/"

    from core.utils import generar_hash_firma_digital, obtener_logos_base64
    firma_digital = generar_hash_firma_digital(solicitud)
    logos = obtener_logos_base64()

    html_string = render_to_string('atencion_ciudadano/constancia_solicitud_pdf.html', {
        'solicitud': solicitud,
        'qr_url': qr_url,
        'firma_digital': firma_digital,
        'fecha_actual': datetime.now(),
        'config': logos['config'],
        'logo_alcaldia_b64': logos['logo_alcaldia'],
        'logo_valera_b64': logos['logo_valera']
    })

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="constancia_solicitud_{solicitud.id}.pdf"'
    pisa_status = pisa.CreatePDF(html_string, dest=response)

    if pisa_status.err:
        return HttpResponse('Error al generar Constancia PDF', status=500)
    return response

@login_required
def configurar_institucion_view(request):
    if not (request.user.is_superuser or (hasattr(request.user, 'perfil') and (request.user.perfil.tiene_permiso('respaldos') or request.user.perfil.tiene_permiso('usuarios')))):
        messages.error(request, "Acceso denegado: Solo administradores o personal autorizado pueden modificar los datos institucionales.")
        return redirect('dashboard')

    from core.models import ConfiguracionInstitucion
    config = ConfiguracionInstitucion.get_solo()

    if request.method == 'POST':
        config.nombre_institucion = request.POST.get('nombre_institucion', config.nombre_institucion).strip()
        config.subtitulo_institucion = request.POST.get('subtitulo_institucion', config.subtitulo_institucion).strip()
        config.rif = request.POST.get('rif', config.rif).strip()
        config.direccion = request.POST.get('direccion', config.direccion).strip()
        config.telefono = request.POST.get('telefono', config.telefono).strip()
        config.email = request.POST.get('email', config.email).strip()

        if request.FILES.get('logo_principal'):
            config.logo_principal = request.FILES['logo_principal']
        if request.FILES.get('logo_secundario'):
            config.logo_secundario = request.FILES['logo_secundario']

        config.save()

        RegistroAuditoria.objects.create(
            usuario=request.user,
            accion="Datos institucionales y logos de membrete actualizados.",
            modulo="Configuración"
        )
        messages.success(request, "¡Datos institucionales y logos actualizados correctamente!")
        return redirect('configurar_institucion')

    return render(request, 'admin_sistema/configuracion_institucion.html', {'config': config})

# --- NUEVOS ENDPOINTS API UX/UI ---
@login_required
def api_busqueda_global(request):
    q = request.GET.get('q', '').strip()
    resultados = []
    if len(q) >= 2:
        # 1. Buscar en Solicitudes de Atención
        solicitudes = SolicitudCiudadano.objects.filter(
            Q(cedula__icontains=q) | Q(nombre_apellido__icontains=q) | Q(descripcion_solicitud__icontains=q)
        )[:5]
        for s in solicitudes:
            resultados.append({
                'tipo': 'solicitud',
                'badge': f'Solicitud #{s.id}',
                'icono': 'bi-envelope-paper-fill text-primary',
                'titulo': f"{s.nombre_apellido} ({s.cedula})",
                'subtitulo': f"{s.tipo_solicitud} - Status: {s.status}",
                'url': reverse('lista_solicitudes_ciudadano') + f"?search={s.cedula}"
            })

        # 2. Buscar en Beneficiarios (Ficha 360)
        beneficiarios = Beneficiario.objects.filter(
            Q(cedula__icontains=q) | Q(nombre_apellido__icontains=q)
        )[:5]
        for b in beneficiarios:
            resultados.append({
                'tipo': 'ciudadano',
                'badge': 'Ficha 360°',
                'icono': 'bi-person-badge-fill text-success',
                'titulo': f"{b.nombre_apellido} - C.I. {b.cedula}",
                'subtitulo': f"Teléfono: {b.telefono or 'Sin registro'}",
                'url': reverse('ficha_ciudadano_360', args=[b.cedula])
            })

        # 3. Buscar en Inventario
        productos = SIA_producto.objects.filter(
            Q(codigo__icontains=q) | Q(descripcion__icontains=q)
        )[:5]
        for p in productos:
            resultados.append({
                'tipo': 'producto',
                'badge': 'Producto',
                'icono': 'bi-box-seam-fill text-warning',
                'titulo': f"{p.descripcion} (Código: {p.codigo})",
                'subtitulo': f"Stock: {p.cantidad} u. | Almacén: {p.almacen or 'General'}",
                'url': reverse('ficha_producto', args=[p.id])
            })

    return JsonResponse({'resultados': resultados})

@login_required
def api_verificar_duplicados_ciudadano(request):
    cedula = request.GET.get('cedula', '').strip()
    if not cedula:
        return JsonResponse({'tiene_pendientes': False, 'solicitudes': []})

    pendientes = SolicitudCiudadano.objects.filter(cedula=cedula).exclude(status__in=['Entregada', 'Rechazada']).order_by('-fecha_solicitud')[:3]
    datos = []
    for p in pendientes:
        datos.append({
            'id': p.id,
            'tipo': p.tipo_solicitud,
            'status': p.status,
            'fecha': p.fecha_solicitud.strftime('%d/%m/%Y')
        })
    return JsonResponse({'tiene_pendientes': len(datos) > 0, 'solicitudes': datos})

@login_required
def api_estadisticas_filtradas(request):
    rango = request.GET.get('rango', 'todos').strip().lower()
    now = timezone.now()
    
    sol_qs = SolicitudCiudadano.objects.all()
    ent_qs = BeneficioEntregado.objects.all()
    
    if rango == 'hoy':
        sol_qs = sol_qs.filter(fecha_solicitud__date=now.date())
        ent_qs = ent_qs.filter(fecha=now.date())
    elif rango == 'semana':
        inicio_semana = now - timezone.timedelta(days=now.weekday())
        sol_qs = sol_qs.filter(fecha_solicitud__gte=inicio_semana)
        ent_qs = ent_qs.filter(fecha__gte=inicio_semana.date())
    elif rango == 'mes':
        sol_qs = sol_qs.filter(fecha_solicitud__month=now.month, fecha_solicitud__year=now.year)
        ent_qs = ent_qs.filter(fecha__month=now.month, fecha__year=now.year)
        
    por_categoria = sol_qs.values('tipo_solicitud').annotate(total=Count('id')).order_by('-total')
    cat_labels = [c['tipo_solicitud'] for c in por_categoria if c['tipo_solicitud']]
    cat_totals = [c['total'] for c in por_categoria if c['tipo_solicitud']]
    
    por_estado = sol_qs.values('status').annotate(total=Count('id'))
    estado_dict = {e['status']: e['total'] for e in por_estado if e['status']}
    cnt_en_evaluacion = estado_dict.get('En Proceso', 0) + estado_dict.get('Recibida', 0) + estado_dict.get('Aprobada', 0)
    
    top_productos = ent_qs.values('descripcion_prod').annotate(total_cant=Sum('cantidad_dada')).order_by('-total_cant')[:6]
    prod_labels = [p['descripcion_prod'] or 'Otros' for p in top_productos]
    prod_totals = [p['total_cant'] for p in top_productos]
    
    total_solicitudes = sol_qs.count()
    total_entregas = ent_qs.count()
    
    return JsonResponse({
        'cat_labels': cat_labels,
        'cat_totals': cat_totals,
        'cnt_en_evaluacion': cnt_en_evaluacion,
        'prod_labels': prod_labels,
        'prod_totals': prod_totals,
        'total_solicitudes': total_solicitudes,
        'total_entregas': total_entregas
    })

@login_required
def api_clear_welcome_splash(request):
    request.session.pop('show_welcome_splash', None)
    return JsonResponse({'status': 'ok'})



